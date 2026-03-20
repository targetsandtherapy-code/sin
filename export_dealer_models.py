"""从小程序网点更新Excel提取经销商-车型对应关系"""
import openpyxl, json

wb = openpyxl.load_workbook(r'D:\api对接\小程序网点更新2-24.xlsx')
ws = wb['一级网点']

headers = [cell.value for cell in ws[1]]
print(f"总列数: {len(headers)}")
print(f"前7列: {headers[:7]}")
print(f"车型列(7起): {headers[7:]}")

model_columns = {}
for i in range(7, len(headers)):
    h = headers[i]
    if h:
        name = str(h).replace('\n', '').strip()
        model_columns[i] = name

print(f"\n识别到 {len(model_columns)} 个车型列:")
for i, name in model_columns.items():
    print(f"  列{i}: {name}")

ws2 = wb['车型配置']
name_to_code = {}
for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, values_only=True):
    code = str(row[1]).strip() if row[1] else ""
    name = str(row[2]).strip() if row[2] else ""
    if code and name:
        name_to_code[name] = code

print(f"\n车型配置表: {len(name_to_code)} 个")
for name, code in name_to_code.items():
    print(f"  {name} -> {code}")

MANUAL_MAP = {
    "艾瑞泽81.6T": "QR76",
    "艾瑞泽82.0T": "QR76",
    "瑞虎7 超能版": "QR64",
    "全新一代瑞虎7PLUS": "QR80",
    "探索06发现版": "QR126",
    "全新艾瑞泽5/艾瑞泽5": "QR42",
    "瑞虎8PRO冠军版": "QR91",
}

col_to_code = {}
unmatched = []
for col_idx, col_name in model_columns.items():
    if col_name in MANUAL_MAP:
        col_to_code[col_idx] = MANUAL_MAP[col_name]
        continue
    matched = False
    for config_name, code in name_to_code.items():
        cn = config_name.replace(' ', '')
        cn2 = col_name.replace(' ', '')
        if cn in cn2 or cn2 in cn:
            col_to_code[col_idx] = code
            matched = True
            break
    if not matched:
        unmatched.append((col_idx, col_name))

print(f"\n已匹配: {len(col_to_code)} 列")
print(f"未匹配: {len(unmatched)} 列")
for idx, name in unmatched:
    print(f"  列{idx}: {name}")

dealer_models = {}
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    erp = str(row[4]) if row[4] else ""
    if not erp:
        continue
    models = []
    for col_idx, code in col_to_code.items():
        if col_idx < len(row) and row[col_idx] and str(row[col_idx]).strip() == '√':
            models.append(code)
    if models:
        dealer_models[erp] = models

print(f"\n有车型数据的经销商: {len(dealer_models)}")
sample = list(dealer_models.items())[:3]
for erp, codes in sample:
    print(f"  ERP {erp}: {codes}")

with open("dealer_models.json", "w", encoding="utf-8") as f:
    json.dump(dealer_models, f, ensure_ascii=False)
print("\nSaved to dealer_models.json")
