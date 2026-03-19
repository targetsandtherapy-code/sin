"""从小程序网点更新Excel导出经销商和车型数据为JSON"""
import openpyxl, json

wb = openpyxl.load_workbook(r'D:\api对接\小程序网点更新2-24.xlsx')

# 经销商
ws = wb['一级网点']
dealers = []
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    if not row[4]:
        continue
    dealers.append({
        "network": str(row[0] or ""),
        "region": str(row[1] or ""),
        "province": str(row[2] or ""),
        "city": str(row[3] or ""),
        "erp_code": str(row[4]),
        "name": str(row[5] or ""),
        "short_name": str(row[6] or ""),
    })
print(f"Dealers: {len(dealers)}")

# 车型
ws2 = wb['车型配置']
models = []
category = ""
for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, values_only=True):
    if row[0]:
        category = str(row[0])
    if row[1]:
        models.append({
            "category": category,
            "code": str(row[1]).strip(),
            "name": str(row[2] or "").strip(),
        })
print(f"Models: {len(models)}")

with open("init_data.json", "w", encoding="utf-8") as f:
    json.dump({"dealers": dealers, "models": models}, f, ensure_ascii=False, indent=2)
print("Saved to init_data.json")
